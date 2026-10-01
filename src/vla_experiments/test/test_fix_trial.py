"""✎ fix trial (morty_gui.fix_trial) on COPIES of a saved trial, the CSVs and
the workbook (skipped where absent): every place the form lives is
corrected, the original data never touched."""
import csv
import json
import shutil
from pathlib import Path

import pytest

EXP = Path.home() / "morty_ws" / "experiments"
SRC = EXP / "experiments_30_09_2026" / "trial_143_LA_S3"


@pytest.fixture
def copies(tmp_path, monkeypatch):
    if not (SRC / "trial.json").exists() or not (EXP / "MMHRI_experiments.xlsx").exists():
        pytest.skip("trial 143 / the workbook not on this machine")
    F = pytest.importorskip("morty_gui.fix_trial")
    d = tmp_path / "experiments_30_09_2026" / SRC.name
    d.mkdir(parents=True)
    for p in SRC.iterdir():
        if p.is_file() and p.stat().st_size < 50_000_000:
            shutil.copy2(p, d / p.name)
    (tmp_path / "mmhri").mkdir()
    for p in (EXP / "mmhri").glob("*.csv"):
        shutil.copy2(p, tmp_path / "mmhri" / p.name)
    shutil.copy2(EXP / "MMHRI_experiments.xlsx", tmp_path / "wb.xlsx")
    monkeypatch.setattr(F, "MMHRI_DIR", str(tmp_path / "mmhri"))
    monkeypatch.setattr(F, "BACKUPS", tmp_path / "backups")
    monkeypatch.setattr(F, "EXPERIMENTS", tmp_path)
    return F, d, tmp_path


def test_fix_corrects_every_place_and_renames(copies):
    F, d, tmp = copies
    before = (SRC / "trial.json").read_bytes()
    lines = F.fix_trial(d, {"Scenario": "S4", "Full Task Success?": "no", "Failure Category / Notes": "t"},
                        xlsx=str(tmp / "wb.xlsx"))
    nd = d.with_name("trial_143_LA_S4")
    assert nd.exists() and not d.exists(), lines
    tj = json.loads((nd / "trial.json").read_text())
    assert tj["Scenario"].startswith("S4:") and tj["Full Task Success?"] == "no" and tj["Corrections"]
    m = json.loads((nd / "meta.json").read_text())
    assert m["outcome"] == "failure" and m["success"] is False and nd.name in m["episode_dir"]
    rows = [r for r in csv.DictReader(open(tmp / "mmhri" / "mmhri_meta.csv")) if r["Trial ID"] == "143"]
    assert rows[-1]["Scenario"].startswith("S4:")
    import openpyxl
    ws = openpyxl.load_workbook(tmp / "wb.xlsx")["Experiments"]
    h = {c.value: i + 1 for i, c in enumerate(ws[1])}
    r = next(r for r in range(2, ws.max_row + 1) if str(ws.cell(r, h["Trial ID"]).value) == "143")
    assert ws.cell(r, h["Scenario ID"]).value == "S4" and ws.cell(r, h["VLA outcome"]).value == "failure"
    assert ws.cell(r, h["Folder"]).value == nd.name
    assert (SRC / "trial.json").read_bytes() == before            # the real trial untouched


def test_no_change_and_dry_run_write_nothing(copies):
    F, d, tmp = copies
    form = F.read_form(d)
    assert F.fix_trial(d, form, xlsx=str(tmp / "wb.xlsx"))[0].endswith("no change")
    out = F.fix_trial(d, {"Scene Layout": "Layout C"}, dry_run=True, xlsx=str(tmp / "wb.xlsx"))
    assert out[-1].startswith("dry run") and d.exists() and not (tmp / "backups").exists()
    with pytest.raises(ValueError):
        F.fix_trial(d, {"Scenario": "S9"}, xlsx=str(tmp / "wb.xlsx"))


def _drop_rows(xlsx, tid):
    """The workbook copy without trial `tid` (as if it was never sent)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    for name, hrow in (("Experiments", 1), ("Stage times", 2)):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        h = {str(c.value).strip(): c.column for c in ws[hrow] if c.value}
        for r in range(ws.max_row, hrow, -1):
            if str(ws.cell(r, h["Trial ID"]).value) == tid:
                ws.delete_rows(r)
    if "_trials" in wb.sheetnames:
        ws = wb["_trials"]
        for r in range(ws.max_row, 0, -1):
            if str(ws.cell(r, 1).value or "").endswith("/trial_143_LA_S3"):
                ws.delete_rows(r)
    wb.save(xlsx)


def test_not_in_excel_is_said_and_the_later_send_carries_the_fix(copies):
    """10-01, trial 188: the layout was corrected before the trial was sent
    — the dry run and the result say so, and sending it afterwards writes
    the corrected layout and folder name."""
    F, d, tmp = copies
    wb = str(tmp / "wb.xlsx")
    assert any(ln == "Excel: its row is corrected too"
               for ln in F.fix_trial(d, {"Scene Layout": "Layout C"}, dry_run=True, xlsx=wb))
    _drop_rows(wb, "143")
    assert any(ln.startswith(F.NOT_IN_EXCEL)
               for ln in F.fix_trial(d, {"Scene Layout": "Layout C"}, dry_run=True, xlsx=wb))
    lines = F.fix_trial(d, {"Scene Layout": "Layout C"}, xlsx=wb)
    nd = d.with_name("trial_143_LC_S3")
    assert nd.exists() and any(ln.startswith(F.NOT_IN_EXCEL) for ln in lines), lines
    res = F.xl.append([str(nd)], wb)
    assert res.get("added") == 1, res
    import openpyxl
    ws = openpyxl.load_workbook(wb)["Experiments"]
    h = {c.value: i + 1 for i, c in enumerate(ws[1])}
    rows = [r for r in range(2, ws.max_row + 1) if str(ws.cell(r, h["Trial ID"]).value) == "143"]
    assert [(ws.cell(r, h["Scene Layout"]).value, ws.cell(r, h["Folder"]).value) for r in rows] == \
        [("Layout C", "trial_143_LC_S3")]


def test_row_found_when_libreoffice_typed_the_cells(copies):
    """A Date cell saved as a date and a Trial ID saved as a number still
    match the trial's '2026-09-30' / '143'."""
    import datetime
    import openpyxl
    F, d, tmp = copies
    wb_path = tmp / "wb.xlsx"
    wb = openpyxl.load_workbook(wb_path)
    ws = wb["Experiments"]
    h = {c.value: i + 1 for i, c in enumerate(ws[1])}
    r = next(r for r in range(2, ws.max_row + 1) if str(ws.cell(r, h["Trial ID"]).value) == "143")
    ws.cell(r, h["Trial ID"]).value = 143.0
    ws.cell(r, h["Date"]).value = datetime.datetime(2026, 9, 30)
    wb.save(wb_path)
    lines = F.fix_trial(d, {"Scene Layout": "Layout C"}, xlsx=str(wb_path))
    assert f"workbook Experiments row {r}" in lines, lines
    ws = openpyxl.load_workbook(wb_path)["Experiments"]
    assert ws.cell(r, h["Scene Layout"]).value == "Layout C"


def test_dialog_offers_to_send_a_trial_without_excel_row(copies, monkeypatch):
    import os
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    QtWidgets = pytest.importorskip("PyQt5.QtWidgets")
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])  # noqa: F841
    D = pytest.importorskip("morty_gui.fix_trial_dialog")
    F, d, tmp = copies
    dlg = D.FixTrialDialog([str(d)], str(d))
    sent = []
    dlg.send_requested.connect(sent.append)
    monkeypatch.setattr(QtWidgets.QMessageBox, "question", staticmethod(lambda *a, **k: QtWidgets.QMessageBox.Yes))
    dlg._on_done(("ok", str(d), ["trial.json", "Excel: its row is corrected too"]))
    assert sent == []
    dlg._on_done(("ok", str(d), ["trial.json", F.NOT_IN_EXCEL + " — nothing to correct there"]))
    assert sent == [str(d)]
