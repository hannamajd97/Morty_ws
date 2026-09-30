"""Move one recorded trial to another scenario — everywhere it is stored.

2026-09-29 (Majd moved 091, 110, 111 and 120 after recording; each was a
hand-held script): the scenario lives in the folder name
(trial_120_LA_S4), trial.json (Scenario, bag / vla_episode paths),
meta.json and viewer.html (paths), mmhri_meta.csv (Scenario) and the Excel
workbook (Experiments: scenario ID + name, folder name and its links; Stage
times: scenario ID; _trials: folder path). This does all of it, with an
optional note (trial.json + mmhri_experiments.csv + Experiments).

Safety: refuses while LibreOffice has the workbook open (the export's
rule); backs up every file it rewrites into
~/morty_ws/backups/relabel_<id>_<time>/; the CSVs are rewritten by
trial_times._update_csv under the trial window's lock (line endings kept),
matched as trial_times matches them (meta: Trial ID + Date, experiments:
Trial ID + Command Used); the workbook's links become absolute file:// URIs
(LibreOffice stores them relative to its real folder on its next save —
relative ones written from experiments/mmhri/ broke when the file moved).

    python3 -m morty_gui.relabel_trial 2026-09-29 120 S5 [--note "..."] [--dry-run]
"""
from __future__ import annotations

import argparse
import csv
import datetime as _dt
import json
import os
import re
import shutil
import sys
from pathlib import Path
from typing import List

from . import excel_export as xl
from .scenario_experiments import SCENARIOS
from .trial_times import MMHRI_DIR, _update_csv, csv_lock

EXPERIMENTS = Path(os.path.expanduser("~/morty_ws/experiments"))
BACKUPS = Path(os.path.expanduser("~/morty_ws/backups"))
NOTE_COL = "Failure Category / Notes"


def _trial_dir(date: str, tid: str) -> Path:
    day = EXPERIMENTS / ("experiments_" + _dt.date.fromisoformat(date).strftime("%d_%m_%Y"))
    hits = sorted(day.glob(f"trial_{tid}_*"))
    if len(hits) != 1:
        raise SystemExit(f"trial {tid} of {date}: {len(hits)} folders in {day}")
    return hits[0]


def relabel(date: str, tid: str, new_sid: str, note: str | None = None,
            dry_run: bool = False, xlsx: str = xl.DEFAULT_XLSX) -> List[str]:
    names = dict(SCENARIOS)
    if new_sid not in names:
        raise SystemExit(f"unknown scenario {new_sid!r} (known: {', '.join(names)})")
    if Path(xlsx).exists() and xl._lock_path(xlsx).exists():
        raise SystemExit("the workbook is open in LibreOffice — close it first (nothing changed)")
    old_dir = _trial_dir(date, tid)
    m = re.fullmatch(r"(trial_\d+_.*_)(S\d+)", old_dir.name)
    if not m:
        raise SystemExit(f"unexpected folder name {old_dir.name!r}")
    old_name, new_name = old_dir.name, m.group(1) + new_sid
    tj = json.loads((old_dir / "trial.json").read_text())
    old_sid = str(tj.get("Scenario", "")).split(":")[0].strip()
    old_full, new_full = str(tj.get("Scenario", "")), f"{new_sid}: {names[new_sid]}"
    cmd = str(tj.get("Command Used", ""))
    done = [f"{old_name} -> {new_name}  ({old_sid} -> {new_sid})"]
    if old_sid == new_sid and note is None:
        return done + ["already there — nothing to do"]
    if dry_run:
        return done + ["dry run — nothing changed"]

    backup = BACKUPS / f"relabel_{tid}_{_dt.datetime.now():%Y%m%d_%H%M%S}"
    backup.mkdir(parents=True)
    for p in [old_dir / "trial.json", old_dir / "meta.json", old_dir / "viewer.html",
              Path(MMHRI_DIR) / "mmhri_meta.csv", Path(MMHRI_DIR) / "mmhri_experiments.csv", Path(xlsx)]:
        if p.exists():
            shutil.copy2(p, backup / p.name)
    (backup / "README.txt").write_text(f"{old_dir} -> {new_name}\n")
    done.append(f"backups in {backup}")

    new_dir = old_dir.with_name(new_name)
    if old_name != new_name:
        if new_dir.exists():
            raise SystemExit(f"{new_dir} exists already")
        old_dir.rename(new_dir)
    # trial.json: the scenario, every path naming the folder, the note
    for k, v in list(tj.items()):
        if isinstance(v, str) and old_name in v:
            tj[k] = v.replace(old_name, new_name)
    tj["Scenario"] = new_full
    if note is not None:
        tj[NOTE_COL] = note
    (new_dir / "trial.json").write_text(json.dumps(tj, indent=2, ensure_ascii=False))
    for f in ("meta.json", "viewer.html"):
        p = new_dir / f
        if p.exists() and old_name in p.read_text():
            p.write_text(p.read_text().replace(old_name, new_name))
    done.append("trial folder: trial.json, meta.json, viewer.html")

    with csv_lock():
        n = _update_csv(os.path.join(MMHRI_DIR, "mmhri_meta.csv"),
                        lambda r: r.get("Trial ID", "").strip() == tid and r.get("Date", "") == date,
                        {"Scenario": (old_full, new_full)})
        done.append(f"mmhri_meta.csv: {n} value(s)")
        if note is not None:
            # _update_csv writes only over the expected old value: read it first
            exp = os.path.join(MMHRI_DIR, "mmhri_experiments.csv")
            match = lambda r: r.get("Trial ID", "").strip() == tid and r.get("Command Used", "") == cmd  # noqa: E731
            with open(exp, newline="") as f:
                rows = [r for r in csv.DictReader(f) if match(r)]
            n = _update_csv(exp, match, {NOTE_COL: (rows[-1].get(NOTE_COL, ""), note)}) if rows else 0
            done.append(f"mmhri_experiments.csv: note ({n} value)")

    if Path(xlsx).exists():
        done += _relabel_workbook(Path(xlsx), date, tid, old_name, new_name, old_sid, new_sid,
                                  names, note)
    return done


def _relabel_workbook(x: Path, date, tid, old_name, new_name, old_sid, new_sid, names, note) -> List[str]:
    import openpyxl
    wb = openpyxl.load_workbook(x)
    out = []
    ws = wb["Experiments"]
    h = xl._header_map(ws, 1)
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, h["Trial ID"]).value == tid and str(ws.cell(r, h["Date"]).value) == date:
            ws.cell(r, h["Scenario ID"]).value = new_sid
            ws.cell(r, h["Scenario"]).value = f"{new_sid}: {names[new_sid]}"
            if note is not None and NOTE_COL in h:
                ws.cell(r, h[NOTE_COL]).value = note
            out.append(f"Experiments row {r}: {old_sid} -> {new_sid}" + (" + note" if note is not None else ""))
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            if cell.value == old_name:
                cell.value = new_name
            if cell.hyperlink and cell.hyperlink.target:
                tgt = fixed = cell.hyperlink.target
                if fixed.startswith("../"):
                    fixed = (EXPERIMENTS / fixed[3:]).as_uri()
                if "/" + old_name in "/" + fixed:
                    fixed = fixed.replace(old_name, new_name)
                    if not fixed.startswith("file://"):
                        fixed = (EXPERIMENTS / fixed).as_uri()
                if fixed != tgt:
                    cell.hyperlink = fixed
    ws = wb["Stage times"]
    h = xl._header_map(ws, 2)
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, h["Trial ID"]).value == tid and str(ws.cell(r, h["Date"]).value) == date:
            ws.cell(r, h["Scenario ID"]).value = new_sid
            out.append(f"Stage times row {r}: {old_sid} -> {new_sid}")
    if "_trials" in wb.sheetnames:
        ws = wb["_trials"]
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if isinstance(v, str) and v.endswith("/" + old_name):
                ws.cell(r, 1).value = v[: -len(old_name)] + new_name
                out.append(f"_trials row {r}")
    tmp = x.with_name(".relabel_" + x.name)
    wb.save(tmp)
    if xl._lock_path(str(x)).exists():
        tmp.unlink()
        return out + ["the workbook was opened meanwhile — NOT written (the rest is done)"]
    os.replace(tmp, x)
    out.append(f"recalc: {xl.recalc(str(x))}")
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("date", help="YYYY-MM-DD")
    ap.add_argument("trial", help="trial id, e.g. 120")
    ap.add_argument("scenario", help="new scenario id, e.g. S5")
    ap.add_argument("--note", default=None, help="Failure Category / Notes")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args(argv)
    for line in relabel(a.date, a.trial.zfill(3), a.scenario.upper(), a.note, a.dry_run):
        print(line)
    return 0


if __name__ == "__main__":
    sys.exit(main())
