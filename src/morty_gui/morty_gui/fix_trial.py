"""Correct a saved trial's form — everywhere it is stored (2026-10-01).

Majd: "when I press Stop & save I forget to change what I wrote here, like
S1 or S2 … or to fill the fields correctly". The form's values go, at ⏹,
into several places; this rewrites all of them from one set of corrections:

  * the folder name (trial_177_LC_S3: layout + scenario) — renamed, and every
    small text file of the folder that names it (trial.json, meta.json,
    recording_manifest.json, bag/postprocess status, viewer.html …);
  * trial.json — the fields, plus a "Corrections" list (when, field, old, new);
  * the VLA episode when "Full Task Success?" changes: meta.json outcome /
    success / provenance and recording_manifest.json — the outcome IS that
    answer (the recorder took it at ⏹), so it follows the correction;
  * mmhri_experiments.csv (the form row) and mmhri_meta.csv (Scenario,
    Modifiers) — under the trial window's CSV lock, each cell only while it
    still holds the old value (trial_times._update_csv);
  * the Excel workbook, if the trial was sent: its Experiments row (by
    header), Scenario ID in Stage times, the folder link / name, the
    _trials path; then the formulas are recalculated.

Safety: refuses while LibreOffice has the workbook open (the export's rule);
backs up every file it rewrites into ~/morty_ws/backups/fix_<id>_<time>/;
the original recording (images, bag, video, steps) is never touched.

    python3 -m morty_gui.fix_trial <trial folder> "Scenario=S4" "Full Task Success?=yes" [--dry-run]
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import os
import re
import shutil
import sys
import time
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from . import excel_export as xl
from .scenario_experiments import SCENARIOS
from .trial_times import MMHRI_DIR, _update_csv, csv_lock

EXPERIMENTS = Path(os.path.expanduser("~/morty_ws/experiments"))
BACKUPS = Path(os.path.expanduser("~/morty_ws/backups"))
YESNO = ["User Intent understood correctly?", "Object Detected?", "Correct Object Selected?",
         "Grasp Success?", "Action Executed Correctly?", "Full Task Success?"]
# the form fields a correction may change, in the form's order
EDITABLE = ["Scene Layout", "Scenario", "Modifiers", "Command Used", "Intended Target",
            "Intended Destination"] + YESNO + ["Failure Category / Notes"]
NAME_RE = re.compile(r"(trial_\d+_)L([A-Z])_(S\d+)")
TEXT_SUFFIXES = (".json", ".html", ".txt", ".srt", ".csv", ".log")


def scenario_full(value: str) -> str:
    """'S4' or 'S4: …' -> the full 'S4: <name>' of scenario_experiments."""
    sid = str(value).split(":")[0].strip().upper()
    names = dict(SCENARIOS)
    if sid not in names:
        raise ValueError(f"unknown scenario {value!r} (known: {', '.join(names)})")
    return f"{sid}: {names[sid]}"


def read_form(folder) -> Dict[str, str]:
    tj = json.loads((Path(folder) / "trial.json").read_text())
    return {k: str(tj.get(k) or "") for k in EDITABLE}


def changes_for(folder, new: Dict[str, str]) -> Dict[str, Tuple[str, str]]:
    """{field: (old, new)} of the fields ``new`` really changes."""
    old = read_form(folder)
    out = {}
    for k, v in new.items():
        if k not in EDITABLE:
            raise ValueError(f"{k!r} is not a form field ({', '.join(EDITABLE)})")
        v = scenario_full(v) if k == "Scenario" else str(v or "").strip()
        if v != old[k].strip():
            out[k] = (old[k], v)
    return out


def folder_name_for(name: str, layout: str, scenario: str) -> str:
    m = NAME_RE.fullmatch(name)
    if not m:
        raise ValueError(f"unexpected folder name {name!r}")
    letter = layout.strip()[-1:].upper() or m.group(2)
    return f"{m.group(1)}L{letter}_{scenario.split(':')[0].strip()}"


def outcome_of(answer: str) -> Tuple[str, Optional[bool]]:
    return {"yes": ("success", True), "no": ("failure", False)}.get(answer.strip().lower(), ("unknown", None))


def set_vla_outcome(folder, answer: str, note: str) -> Optional[str]:
    """The VLA episode in ``folder`` takes the human verdict ``answer``
    ("Full Task Success?": yes / no / empty): meta.json outcome, success,
    provenance + a correction record, and the manifest's outcome. Returns
    "old -> new", or None (no episode, or nothing to change). Used by
    fix_trial and by the trial window's 💾 Save (2026-10-01)."""
    folder = Path(folder)
    mp = folder / "meta.json"
    if not mp.exists():
        return None
    outcome, success = outcome_of(answer)
    m = json.loads(mp.read_text())
    was = m.get("outcome")
    if was == outcome:
        return None
    m["outcome"], m["success"] = outcome, success
    m["outcome_provenance"] = m["note"] = note
    m.setdefault("outcome_corrections", []).append(
        {"at": time.time(), "was": was, "now": outcome, "answer": answer})
    _atomic_write(mp, json.dumps(m, indent=2, default=str))
    man_p = folder / "recording_manifest.json"
    if man_p.exists():
        man = json.loads(man_p.read_text())
        man["outcome"] = outcome
        _atomic_write(man_p, json.dumps(man, indent=2))
    return f"{was} -> {outcome}"


def _atomic_write(p: Path, text: str) -> None:
    tmp = p.with_name("." + p.name + ".fix_tmp")
    tmp.write_text(text)
    os.replace(tmp, p)


def fix_trial(folder, new: Dict[str, str], dry_run: bool = False, xlsx: str = xl.DEFAULT_XLSX) -> List[str]:
    """Apply ``new`` ({form field: value}) to the saved trial in ``folder``.
    Returns what was done, line by line (the dry run: what would be)."""
    folder = Path(folder)
    if not (folder / "trial.json").exists():
        raise ValueError(f"{folder} has no trial.json (not a saved trial)")
    ch = changes_for(folder, new)
    tj = json.loads((folder / "trial.json").read_text())
    tid, date = str(tj.get("Trial ID", "")).strip(), str(tj.get("Date", ""))
    layout = ch.get("Scene Layout", (None, tj.get("Scene Layout", "")))[1]
    scenario = ch.get("Scenario", (None, tj.get("Scenario", "")))[1]
    old_name = folder.name
    new_name = folder_name_for(old_name, layout, scenario) if ("Scene Layout" in ch or "Scenario" in ch) else old_name
    out = [f"trial {tid} of {date}: " + ("; ".join(f"{k}: {o or '(empty)'!r} -> {n or '(empty)'!r}"
                                                  for k, (o, n) in ch.items()) or "no change")]
    if new_name != old_name:
        out.append(f"folder {old_name} -> {new_name}")
    if not ch:
        return out
    if Path(xlsx).exists() and xl._lock_path(xlsx).exists():
        raise RuntimeError("the workbook is open in LibreOffice — close it first (nothing changed)")
    new_dir = folder.with_name(new_name)
    if new_name != old_name and new_dir.exists():
        raise RuntimeError(f"{new_dir} exists already (nothing changed)")
    ep_meta = folder / "meta.json" if (folder / "meta.json").exists() else None
    if dry_run:
        if "Full Task Success?" in ch and ep_meta:
            out.append(f"VLA episode outcome -> {outcome_of(ch['Full Task Success?'][1])[0]}")
        return out + ["dry run — nothing changed"]

    # 1. backups of everything rewritten below
    backup = BACKUPS / f"fix_{tid}_{_dt.datetime.now():%Y%m%d_%H%M%S}"
    backup.mkdir(parents=True)
    texts = [p for p in folder.iterdir()
             if p.is_file() and p.suffix in TEXT_SUFFIXES and p.stat().st_size < 50_000_000]
    for p in texts + [Path(MMHRI_DIR) / "mmhri_meta.csv", Path(MMHRI_DIR) / "mmhri_experiments.csv", Path(xlsx)]:
        if p.exists():
            shutil.copy2(p, backup / p.name)
    (backup / "README.txt").write_text(f"{folder}\n{json.dumps(ch, indent=2, ensure_ascii=False)}\n")
    out.append(f"backups in {backup}")

    # 2. the folder and the files naming it
    if new_name != old_name:
        folder.rename(new_dir)
        n = 0
        for p in texts:
            q = new_dir / p.name
            s = q.read_text(errors="surrogateescape")
            if old_name in s:
                _atomic_write(q, s.replace(old_name, new_name))
                n += 1
        out.append(f"folder renamed; {n} file(s) inside updated to the new name")
        tj = json.loads((new_dir / "trial.json").read_text())

    # 3. trial.json
    for k, (_o, v) in ch.items():
        tj[k] = v
    tj.setdefault("Corrections", []).append(
        {"at": _dt.datetime.now().isoformat(timespec="seconds"),
         "changes": {k: {"old": o, "new": v} for k, (o, v) in ch.items()}})
    _atomic_write(new_dir / "trial.json", json.dumps(tj, indent=2, ensure_ascii=False))
    out.append("trial.json")

    # 4. the VLA episode's outcome follows "Full Task Success?"
    if "Full Task Success?" in ch:
        done = set_vla_outcome(new_dir, ch["Full Task Success?"][1], "outcome: human (corrected in the trial window)")
        if done:
            out.append(f"VLA episode outcome {done}")

    # 5. the CSVs (form row: Trial ID + the OLD command; meta row: Trial ID + Date)
    cmd_old = ch.get("Command Used", (tj.get("Command Used", ""), None))[0]
    with csv_lock():
        exp_cols = [k for k in ch if k not in ("Scenario",)]
        n1 = _update_csv(os.path.join(MMHRI_DIR, "mmhri_experiments.csv"),
                         lambda r: r.get("Trial ID", "").strip() == tid and r.get("Command Used", "") == cmd_old,
                         {k: ch[k] for k in exp_cols})
        n2 = _update_csv(os.path.join(MMHRI_DIR, "mmhri_meta.csv"),
                         lambda r: r.get("Trial ID", "").strip() == tid and r.get("Date", "") == date,
                         {k: ch[k] for k in ("Scenario", "Modifiers") if k in ch})
    out.append(f"mmhri_experiments.csv: {n1} value(s); mmhri_meta.csv: {n2} value(s)")

    # 6. the workbook, if the trial is in it
    if Path(xlsx).exists():
        out += _fix_workbook(Path(xlsx), tid, date, old_name, new_name, ch,
                             outcome_of(ch["Full Task Success?"][1])[0]
                             if "Full Task Success?" in ch and (new_dir / "meta.json").exists() else None)
    return out


def _fix_workbook(x: Path, tid, date, old_name, new_name, ch, vla_outcome) -> List[str]:
    import openpyxl
    wb = openpyxl.load_workbook(x)
    out, touched = [], False
    ws = wb["Experiments"]
    h = xl._header_map(ws, 1)
    for r in range(2, ws.max_row + 1):
        if not (str(ws.cell(r, h["Trial ID"]).value) == tid and str(ws.cell(r, h["Date"]).value) == date):
            continue
        for k, (_o, v) in ch.items():
            if k == "Scenario":
                if "Scenario ID" in h:
                    ws.cell(r, h["Scenario ID"]).value = v.split(":")[0].strip()
                if "Scenario" in h:
                    ws.cell(r, h["Scenario"]).value = v
            elif k in h:
                ws.cell(r, h[k]).value = v or None
        if vla_outcome is not None and "VLA outcome" in h:
            ws.cell(r, h["VLA outcome"]).value = vla_outcome
        if new_name != old_name:
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(r, c)
                if cell.value == old_name:
                    cell.value = new_name
                if cell.hyperlink and cell.hyperlink.target:
                    tgt = fixed = cell.hyperlink.target
                    if fixed.startswith("../"):
                        fixed = (EXPERIMENTS / fixed[3:]).as_uri()
                    if "/" + old_name in "/" + fixed:
                        fixed = fixed.replace(old_name, new_name)
                        if not fixed.startswith("file://"):
                            fixed = (EXPERIMENTS / fixed).as_uri()
                    if fixed != tgt:
                        cell.hyperlink = fixed
        out.append(f"workbook Experiments row {r}")
        touched = True
    if "Scenario" in ch and "Stage times" in wb.sheetnames:
        ws = wb["Stage times"]
        h = xl._header_map(ws, 2)
        for r in range(3, ws.max_row + 1):
            if str(ws.cell(r, h["Trial ID"]).value) == tid and str(ws.cell(r, h["Date"]).value) == date:
                ws.cell(r, h["Scenario ID"]).value = ch["Scenario"][1].split(":")[0].strip()
                out.append(f"workbook Stage times row {r}")
                touched = True
    if new_name != old_name and "_trials" in wb.sheetnames:
        ws = wb["_trials"]
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if isinstance(v, str) and v.endswith("/" + old_name):
                ws.cell(r, 1).value = v[: -len(old_name)] + new_name
                touched = True
    if not touched:
        return ["workbook: the trial is not in it (send it with 📊 when ready)"]
    tmp = x.with_name(".fix_" + x.name)
    wb.save(tmp)
    if xl._lock_path(str(x)).exists():
        tmp.unlink()
        return out + ["the workbook was opened meanwhile — NOT written (the rest is done)"]
    os.replace(tmp, x)
    out.append(f"recalc: {xl.recalc(str(x))}")
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("folder", help="the trial folder")
    ap.add_argument("changes", nargs="+", help='"Field=value", e.g. "Scenario=S4" "Full Task Success?=yes"')
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args(argv)
    new = dict(c.split("=", 1) for c in a.changes)
    for line in fix_trial(Path(a.folder).expanduser(), new, a.dry_run):
        print(line)
    return 0


if __name__ == "__main__":
    sys.exit(main())
