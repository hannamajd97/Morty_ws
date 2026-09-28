"""The MMHRI experiments as one Excel workbook (2026-09-28, Majd: "a button in
the experiments window that sends the whole recorded stuff to a sheet, with a
selection of the experiments to send, updated after every experiment").

    python3 -m morty_gui.excel_export [--all | <trial dir> ...] [--out file.xlsx]

The window's "📊 Excel export" does the same through ExcelExportWindow. The
workbook (default experiments/mmhri/MMHRI_experiments.xlsx) is REBUILT from
the trial folders on every send — the folders stay the one source of truth
(trial.json, timeline.json, meta.json, conversation.txt), so a trial whose
times were finalized after an earlier send is up to date after the next.
Only the sheets written here are replaced; a sheet you add yourself is kept.

Sheets:
  Experiments       one row per trial: the doc's table (Trial ID … Failure
                    Category / Notes) plus date, scenario, level, modifiers,
                    how the pick was started, the key times, the VLA task and
                    outcome, and links to the folder, the viewer and the video
  Stage times       every stage time of the trial (see Definitions)
  Conversation      what was said and commanded, one line per event, with
                    the time in the trial
  VLA instructions  the episode's instructions: start time, text, what was
                    said, the steps they cover
  Summary           per scenario: trials, verdicts, success rate and mean
                    times — Excel formulas over the Experiments sheet
  Definitions       what every column means
  _trials           (hidden) the trial folders in this workbook — the
                    window reads it to tick what is already in Excel
"""
from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional

EXPERIMENTS_DIR = os.path.expanduser("~/morty_ws/experiments")
MMHRI_DIR = os.path.join(EXPERIMENTS_DIR, "mmhri")
DEFAULT_XLSX = os.path.join(MMHRI_DIR, "MMHRI_experiments.xlsx")

GENERATED = ["Experiments", "Stage times", "Conversation", "VLA instructions", "Summary", "Definitions", "_trials"]
VERDICTS = ["User Intent understood correctly?", "Object Detected?", "Correct Object Selected?",
            "Grasp Success?", "Action Executed Correctly?", "Full Task Success?"]

# (header, key in the trial record, number format or None, width)
EXP_COLUMNS = [
    ("Trial ID", "trial_id", "@", 8), ("Date", "date", None, 11), ("Task ID", "task_id", None, 8),
    ("Scene Layout", "layout", None, 10), ("Scenario ID", "scenario_id", None, 9), ("Scenario", "scenario", None, 34),
    ("Level", "level", None, 18), ("Modifiers", "modifiers", None, 12), ("Command Used", "command", None, 36),
    ("Intended Target", "target", None, 16), ("Intended Destination", "destination", None, 14),
] + [(v, "v:" + v, None, 12) for v in VERDICTS] + [
    ("Interaction Time (s)", "interaction_col_s", "0.00", 11), ("Total Time (s)", "total_time_s", "0.00", 10),
    ("Failure Category / Notes", "notes", None, 30),
    ("Started by", "started_by", None, 11),
    ("Understood by", "understood_by", None, 13),
    ("GraspGen in background?", "graspgen_background", None, 11),
    ("Execution plan", "planner", None, 10),
    ("Plan checked in background?", "plan_background", None, 11),
    ("Grasp executed", "grasp_used", None, 10),
    ("Command → object handed over (s)", "release_s", "0.00", 12),
    ("Waiting for your answers (s)", "human_wait_s", "0.00", 11),
    ("Robot time (s)", "robot_time_s", "0.00", 10),
    ("Times", "times_state", None, 10),
    ("VLA task", "vla_task", None, 30), ("VLA outcome", "vla_outcome", None, 10),
    ("Folder", "link_folder", None, 26), ("Viewer", "link_viewer", None, 10), ("Video", "link_video", None, 10),
]

# (group, header, key) — the Stage times sheet, in the order things happen
STAGE_COLUMNS = [
    ("Interaction", "Interaction Time: command given → arm starts moving (s)", "interaction_s"),
    ("Voice", "End of your voice → transcript (s)", "voice_to_transcript_s"),
    ("Voice", "You stopped speaking → command (s)", "speech_end_to_intent_s"),
    ("Voice", "Understanding: transcript → command (s)", "intent_s"),
    ("Voice", "Understood by", "understood_by"),
    ("Choice", "Command → object chosen, incl. clarification (s)", "choice_s"),
    ("Perception", "Object chosen → grasp candidates (s)", "yolo_s"),
    ("GraspGen", "GraspGen in background?", "graspgen_background"),
    ("GraspGen", "GraspGen compute (s)", "graspgen_compute_s"),
    ("GraspGen", "GraspGen waited for after the choice (s)", "graspgen_wait_s"),
    ("GraspGen", "GraspGen ready before the choice (s)", "graspgen_lead_s"),
    ("Execution plan", "Execution plan", "planner"),
    ("Execution plan", "Plan checked in background?", "plan_background"),
    ("Execution plan", "Plan compute (s)", "plan_compute_s"),
    ("Execution plan", "Candidates → motion start (s)", "planning_s"),
    ("Execution plan", "Grasp executed", "grasp_used"),
    ("Execution plan", "Grasps checked", "grasp_probes"),
    ("Motion", "Approach (s)", "approach_s"),
    ("Motion", "Arrival → grasp verdict (s)", "grasp_s"),
    ("Motion", "Waiting for the drop side (s)", "wait_destination_s"),
    ("Motion", "Transport to the drop (s)", "transport_s"),
    ("Motion", "Motion start → home (s)", "execution_s"),
    ("Totals", "Command → object handed over (s)", "release_s"),
    ("Totals", "Waiting for your answers (s)", "human_wait_s"),
    ("Totals", "Robot time (s)", "robot_time_s"),
    ("Totals", "Total (s)", "total_s"),
    ("Dialogue", "Morty's questions", "morty_questions"),
    ("Dialogue", "Your sentences", "your_utterances"),
]
STAGE_TEXT = {"understood_by", "graspgen_background", "plan_background", "grasp_used", "planner", "path"}

DEFINITIONS = [
    ("Where the data comes from", "Every sheet is rebuilt from the trial folders (trial.json, timeline.json, meta.json) each time "
     "the window sends. Edit a trial in the MMHRI window (or its trial.json), not in these sheets — your edits here would be "
     "replaced. Sheets you add yourself are kept."),
    ("Interaction Time (s)", "The experiment doc's definition: from the moment the command was GIVEN (you stopped saying it; for "
     "a pick started by pointing, the moment the object was chosen) until the arm STARTS MOVING. Includes speech recognition, "
     "understanding, any clarification ('which one?'), choosing the object and planning."),
    ("Total Time (s)", "From the first command (or the pointing choice) to the end of the robot's cycle (back home)."),
    ("Started by", "voice = the task began with a spoken command; pointing/click = the object was chosen by pointing or a "
     "click before anything was said."),
    ("End of your voice → transcript (s)", "The end-of-speech wait (0.5-0.8 s, to be sure you finished) plus speech recognition."),
    ("You stopped speaking → command (s)", "The end of the operator's voice → the robot command for that sentence."),
    ("Understanding: transcript → command (s)", "The accepted transcript that caused the command → the command."),
    ("Understood by", "instant rule = the fixed phrase layer (about 0.01 s); language model = Claude (about 1 s)."),
    ("Command → object chosen, incl. clarification (s)", "The first command → the pipeline committing the object. Includes a "
     "clarification dialogue ('which one?') or the pointing, and handing the object's name to the pipeline."),
    ("Object chosen → grasp candidates (s)", "The committed choice → grasp candidates published (capture, segmentation, the "
     "stored-grasp lookup, or GraspGen when nothing was stored)."),
    ("GraspGen in background?", "yes = GraspGen had already computed this object's grasps before the command (precompute); "
     "no = it ran after the command, and the operator waited for it."),
    ("GraspGen compute (s)", "Duration of the GraspGen run whose grasps the robot executed."),
    ("GraspGen waited for after the choice (s)", "How much of that run the operator actually waited for (0 = ready before)."),
    ("GraspGen ready before the choice (s)", "How long before the choice the grasps were ready."),
    ("Execution plan", "How the approach was planned: curve = the smooth approach path (bend + straight slide); lines = "
     "blended straight-line segments, the fallback."),
    ("Plan checked in background?", "yes = the executor had already verified this grasp could be reached (precheck hit), so "
     "it moved at once; no = it checked after the candidates arrived (precheck miss) and the operator waited for it."),
    ("Plan compute (s)", "Time to verify the plan: the background probe (when checked in the background) or the probe after "
     "the candidates arrived."),
    ("Candidates → motion start (s)", "What the operator actually waited in the executor, from the candidates to moving."),
    ("Grasp executed", "Which of the GraspGen candidates was executed ('1 of 26' = the best-ranked one)."),
    ("Grasps checked", "How many candidates the executor probed before executing one."),
    ("Command → object handed over (s)", "The command given → the gripper opened at the destination."),
    ("Waiting for your answers (s)", "Time Morty waited for the operator: at 'Where should I drop it?', 'Should I open the "
     "gripper?' and 'Should I go back home?'."),
    ("Robot time (s)", "Total time minus the time Morty waited for the operator."),
    ("Times", "final = rebuilt from exact event times after the trial (the messages' own stamps, else the earliest recorder); "
     "preliminary = the window's own times at ⏹ (can be late under load)."),
    ("Summary: success rate", "Full Task Success = yes divided by (yes + no); trials without a verdict are not counted."),
]


def _num(v):
    if v is None or v == "":
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if f == f else None


def _load_json(p: Path):
    try:
        return json.loads(p.read_text())
    except Exception:
        return None


def all_trial_dirs() -> List[Path]:
    """Every trial folder with a saved trial.json (the window's layouts)."""
    out = set()
    base = Path(EXPERIMENTS_DIR)

    def scan(d: Path, depth: int) -> None:
        try:
            for p in sorted(d.iterdir()):
                if not p.is_dir():
                    continue
                if p.name.startswith("trial_"):
                    if (p / "trial.json").exists():
                        out.add(p)
                elif depth > 0 and not p.name.startswith("."):
                    scan(p, depth - 1)
        except OSError:
            pass
    scan(base, 2)
    return sorted(out)


def _sort_key(rec: dict):
    return (rec.get("date") or "", rec.get("trial_id") or "", rec.get("folder") or "")


def trial_record(d: Path) -> Optional[dict]:
    """Everything the workbook shows about one trial."""
    d = Path(d)
    tj = _load_json(d / "trial.json")
    if not isinstance(tj, dict):
        return None
    st = tj.get("Stage Times") if isinstance(tj.get("Stage Times"), dict) else {}
    scen = str(tj.get("Scenario") or "")
    rec = {
        "folder": str(d), "trial_id": str(tj.get("Trial ID") or ""), "date": str(tj.get("Date") or ""),
        "task_id": tj.get("Task ID") or None, "layout": tj.get("Scene Layout") or None,
        "scenario_id": scen.split(":")[0].strip() or None, "scenario": scen or None, "level": tj.get("Level") or None,
        "modifiers": tj.get("Modifiers") or None, "command": tj.get("Command Used") or None,
        "target": tj.get("Intended Target") or None, "destination": tj.get("Intended Destination") or None,
        "interaction_col_s": _num(tj.get("Interaction Time (s)")), "total_time_s": _num(tj.get("Total Time (s)")),
        "notes": tj.get("Failure Category / Notes") or None,
        "started_by": {"voice": "voice", "choice": "pointing/click"}.get(st.get("started_by"), None),
        "times_state": "final" if st.get("event_times") else ("preliminary" if st else None),
        "gui_late_max_s": _num(((st.get("event_times") or {}).get("gui_late_by_s") or {}).get("max")),
        "path": st.get("path"),
    }
    for v in VERDICTS:
        val = str(tj.get(v) or "").strip()
        rec["v:" + v] = val or None
    for _g, _h, k in STAGE_COLUMNS:
        rec[k] = (st.get(k) or None) if k in STAGE_TEXT else _num(st.get(k))
    meta = _load_json(d / "meta.json") if (d / "meta.json").exists() else None
    rec["vla_task"] = (meta or {}).get("task") if meta else None
    rec["vla_outcome"] = (meta or {}).get("outcome") if meta else None
    rec["link_folder"] = d.as_uri()
    rec["link_viewer"] = (d / "viewer.html").as_uri() if (d / "viewer.html").exists() else None
    video = d / "trial.mp4" if (d / "trial.mp4").exists() else (d / "camera.mp4" if (d / "camera.mp4").exists() else None)
    rec["link_video"] = video.as_uri() if video else None
    # conversation: exact times from the (finalized) timeline, else conversation.txt
    conv = []
    tl = _load_json(d / "timeline.json") if (d / "timeline.json").exists() else None
    if isinstance(tl, dict) and tl.get("events"):
        for e in tl["events"]:
            src, txt = e.get("source"), str(e.get("text") or "")
            if src == "transcript":
                conv.append((_num(e.get("t")), "You", txt))
            elif src == "morty_said":
                conv.append((_num(e.get("t")), "Morty", txt))
            elif src == "intent":
                try:
                    j = json.loads(txt)
                    a = j.get("attributes") or j.get("object_attributes") or {}
                    extra = " ".join(f"{k}={a[k]}" for k in ("position", "colour", "destination", "then")
                                     if isinstance(a, dict) and a.get(k))
                    txt = " ".join(x for x in (str(j.get("action_intent") or ""), str(j.get("object_intent") or "").replace("_", " "), extra) if x and x != "None")
                except Exception:
                    pass
                conv.append((_num(e.get("t")), "Robot command", txt))
    elif (d / "conversation.txt").exists():
        for line in (d / "conversation.txt").read_text(errors="ignore").splitlines():
            m = re.match(r"^\[\s*([\d.]+)s\]\s+([^:]+):\s*(.*)$", line)
            if m:
                who = "You" if m.group(2).startswith("You") else "Morty"
                conv.append((_num(m.group(1)), who, m.group(3)))
    rec["conversation"] = conv
    # VLA instructions with the steps they cover
    ins = []
    if meta and meta.get("instructions"):
        spans: Dict[int, tuple] = {}
        try:
            import numpy as np
            with np.load(d / "steps.npz") as z:
                if "instruction_idx" in z.files:
                    ii = z["instruction_idx"]
                    for k in range(len(meta["instructions"])):
                        on = np.where(ii == k)[0]
                        if len(on):
                            spans[k] = (int(on[0]), int(on[-1]))
        except Exception:
            pass
        for k, i in enumerate(meta["instructions"]):
            attrs = i.get("attributes") if isinstance(i.get("attributes"), dict) else {}
            ins.append({"n": k + 1, "t": _num(i.get("t")), "text": i.get("instruction"),
                        "said": i.get("transcript") or (f"({attrs.get('source')})" if attrs.get("source") else None),
                        "action": i.get("action_intent"),
                        "steps": f"{spans[k][0]}–{spans[k][1]}" if k in spans else None,
                        "superseded": "yes" if i.get("superseded") else None})
    rec["instructions"] = ins
    return rec


def exported_folders(path: str = DEFAULT_XLSX) -> List[str]:
    """The trial folders the workbook holds (its hidden _trials sheet)."""
    if not os.path.exists(path):
        return []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        if "_trials" not in wb.sheetnames:
            return []
        out = [str(r[0]) for r in wb["_trials"].iter_rows(min_row=2, values_only=True) if r and r[0]]
        wb.close()
        return out
    except Exception:
        return []


def build(folders: List[str], path: str = DEFAULT_XLSX) -> dict:
    """(Re)write the generated sheets for exactly these trial folders."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties  # noqa: F401 (keeps openpyxl import checks honest)

    recs = [r for r in (trial_record(Path(f)) for f in folders) if r]
    recs.sort(key=_sort_key)
    if os.path.exists(path):
        wb = load_workbook(path)
        for name in GENERATED:
            if name in wb.sheetnames:
                del wb[name]
    else:
        wb = Workbook()
        wb.remove(wb.active)
    FONT = Font(name="Arial", size=10)
    BOLD = Font(name="Arial", size=10, bold=True)
    HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
    TITLE = Font(name="Arial", size=12, bold=True)
    thin = Side(style="thin", color="B7B7B7")
    BORDER = Border(bottom=thin)
    WRAP = Alignment(wrap_text=True, vertical="top")
    TOP = Alignment(vertical="top")

    def header(ws, row, heads, widths=None, notes=None):
        for c, h in enumerate(heads, 1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font = BOLD
            cell.fill = HEAD_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = BORDER
            if notes and notes.get(h):
                cell.comment = Comment(notes[h], "MMHRI window")
            if widths:
                ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
        ws.row_dimensions[row].height = 42

    defs = dict(DEFINITIONS)
    pos = len(GENERATED)
    # ── Experiments ─────────────────────────────────────────────────────
    ws = wb.create_sheet("Experiments", 0)
    header(ws, 1, [h for h, _k, _f, _w in EXP_COLUMNS], [w for *_x, w in EXP_COLUMNS], defs)
    for r, rec in enumerate(recs, 2):
        for c, (h, k, fmt, _w) in enumerate(EXP_COLUMNS, 1):
            v = rec.get(k)
            cell = ws.cell(row=r, column=c)
            if k.startswith("link_"):
                if v:
                    cell.value = {"link_folder": os.path.basename(rec["folder"]), "link_viewer": "viewer", "link_video": "video"}[k]
                    cell.hyperlink = v
                    cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
                continue
            cell.value = v
            cell.font = FONT
            cell.alignment = TOP
            if fmt and v is not None:
                cell.number_format = fmt
    n = len(recs)
    last = n + 1
    ws.freeze_panes = "C2"
    if n:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(EXP_COLUMNS))}{last}"
        green = PatternFill("solid", fgColor="E2F0D9")
        red = PatternFill("solid", fgColor="F8D7DA")
        c0 = [h for h, *_x in EXP_COLUMNS].index(VERDICTS[0]) + 1
        rng = f"{get_column_letter(c0)}2:{get_column_letter(c0 + len(VERDICTS) - 1)}{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"yes"'], fill=green))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"no"'], fill=red))
    # ── Stage times ─────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Stage times", 1)
    heads = ["Trial ID", "Date", "Scenario ID", "Started by"] + [h for _g, h, _k in STAGE_COLUMNS] + \
            ["Path", "Times", "Window was late by, max (s)"]
    groups = ["", "", "", ""] + [g for g, _h, _k in STAGE_COLUMNS] + ["", "", ""]
    prev = None
    for c, g in enumerate(groups, 1):          # a group row above the headers
        if g and g != prev:
            gc = ws2.cell(row=1, column=c, value=g)
            gc.font = BOLD
            gc.fill = PatternFill("solid", fgColor="B8CCE4")
        elif g:
            ws2.cell(row=1, column=c).fill = PatternFill("solid", fgColor="B8CCE4")
        prev = g
    header(ws2, 2, heads, [8, 11, 9, 12] + [12] * len(STAGE_COLUMNS) + [12, 10, 12], defs)
    for r, rec in enumerate(recs, 3):
        vals = [rec["trial_id"], rec["date"], rec["scenario_id"], rec["started_by"]] + \
               [rec.get(k) for _g, _h, k in STAGE_COLUMNS] + [rec.get("path"), rec.get("times_state"), rec.get("gui_late_max_s")]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(row=r, column=c, value=v)
            cell.font = FONT
            if isinstance(v, float):
                cell.number_format = "0.00"
    ws2.freeze_panes = "E3"
    if n:
        ws2.auto_filter.ref = f"A2:{get_column_letter(len(heads))}{last + 1}"
    # ── Conversation ────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Conversation", 2)
    header(ws3, 1, ["Trial ID", "Date", "t in trial (s)", "Who", "Text"], [8, 11, 11, 14, 90])
    r = 2
    for rec in recs:
        for t, who, txt in rec["conversation"]:
            for c, v in enumerate([rec["trial_id"], rec["date"], t, who, txt], 1):
                cell = ws3.cell(row=r, column=c, value=v)
                cell.font = BOLD if (c == 4 and who == "You") else FONT
                cell.alignment = WRAP if c == 5 else TOP
                if c == 3 and v is not None:
                    cell.number_format = "0.00"
            r += 1
    ws3.freeze_panes = "A2"
    if r > 2:
        ws3.auto_filter.ref = f"A1:E{r - 1}"
    # ── VLA instructions ────────────────────────────────────────────────
    ws4 = wb.create_sheet("VLA instructions", 3)
    header(ws4, 1, ["Trial ID", "Date", "#", "Starts at t (s)", "Instruction", "Said / from", "Action", "Steps", "Superseded"],
           [8, 11, 5, 11, 46, 46, 22, 11, 11])
    r = 2
    for rec in recs:
        for i in rec["instructions"]:
            for c, v in enumerate([rec["trial_id"], rec["date"], i["n"], i["t"], i["text"], i["said"], i["action"], i["steps"], i["superseded"]], 1):
                cell = ws4.cell(row=r, column=c, value=v)
                cell.font = FONT
                cell.alignment = WRAP if c in (5, 6) else TOP
                if c == 4 and v is not None:
                    cell.number_format = "0.00"
            r += 1
    ws4.freeze_panes = "A2"
    if r > 2:
        ws4.auto_filter.ref = f"A1:I{r - 1}"
    # ── Summary (formulas over Experiments) ─────────────────────────────
    ws5 = wb.create_sheet("Summary", 4)
    ws5["A1"] = "MMHRI experiments — summary per scenario"
    ws5["A1"].font = TITLE
    ws5["A2"] = (f"Formulas over the Experiments sheet ({n} trials). Success rate = Full Task Success yes / (yes + no); "
                 "averages ignore empty cells.")
    ws5["A2"].font = Font(name="Arial", size=9, italic=True, color="595959")
    col = {h: get_column_letter(i) for i, (h, *_x) in enumerate(EXP_COLUMNS, 1)}

    def rng(h):
        return f"Experiments!${col[h]}$2:${col[h]}${max(last, 2)}"

    sheads = ["Scenario ID", "Scenario", "Trials", "Full task success: yes", "Full task success: no", "Success rate",
              "Grasp success: yes", "Avg Interaction Time (s)", "Avg total time (s)", "Avg robot time (s)",
              "Avg waiting for your answers (s)"]
    header(ws5, 4, sheads, [10, 44, 8, 12, 12, 10, 12, 12, 14, 16, 13], defs)
    from .scenario_experiments import SCENARIOS
    first = 5
    for i, (sid, txt) in enumerate(SCENARIOS):
        rr = first + i
        ws5.cell(row=rr, column=1, value=sid)
        ws5.cell(row=rr, column=2, value=txt)
        ws5.cell(row=rr, column=3, value=f"=COUNTIFS({rng('Scenario ID')},$A{rr})")
        ws5.cell(row=rr, column=4, value=f'=COUNTIFS({rng("Scenario ID")},$A{rr},{rng("Full Task Success?")},"yes")')
        ws5.cell(row=rr, column=5, value=f'=COUNTIFS({rng("Scenario ID")},$A{rr},{rng("Full Task Success?")},"no")')
        ws5.cell(row=rr, column=6, value=f'=IF(D{rr}+E{rr}=0,"",D{rr}/(D{rr}+E{rr}))')
        ws5.cell(row=rr, column=7, value=f'=COUNTIFS({rng("Scenario ID")},$A{rr},{rng("Grasp Success?")},"yes")')
        for c, h in ((8, "Interaction Time (s)"), (9, "Total Time (s)"), (10, "Robot time (s)"),
                     (11, "Waiting for your answers (s)")):
            ws5.cell(row=rr, column=c, value=f'=IF(COUNTIFS({rng("Scenario ID")},$A{rr},{rng(h)},"<>")=0,"",'
                                                f'AVERAGEIFS({rng(h)},{rng("Scenario ID")},$A{rr}))')
    tr = first + len(SCENARIOS)
    ws5.cell(row=tr, column=1, value="All")
    for c in (3, 4, 5, 7):
        L = get_column_letter(c)
        ws5.cell(row=tr, column=c, value=f"=SUM({L}{first}:{L}{tr - 1})")
    ws5.cell(row=tr, column=6, value=f'=IF(D{tr}+E{tr}=0,"",D{tr}/(D{tr}+E{tr}))')
    for c, h in ((8, "Interaction Time (s)"), (9, "Total Time (s)"), (10, "Robot time (s)"),
                 (11, "Waiting for your answers (s)")):
        ws5.cell(row=tr, column=c, value=f'=IF(COUNT({rng(h)})=0,"",AVERAGE({rng(h)}))')
    for rr in range(first, tr + 1):
        for c in range(1, len(sheads) + 1):
            cell = ws5.cell(row=rr, column=c)
            cell.font = BOLD if rr == tr else FONT
            cell.alignment = TOP
            if c == 6:
                cell.number_format = "0%"
            elif c >= 8:
                cell.number_format = "0.00"
    # trials per scenario x layout (the window's progress grid)
    gr = tr + 3
    ws5.cell(row=gr - 1, column=1, value="Trials per scenario and layout").font = BOLD
    layouts = sorted({rec["layout"] for rec in recs if rec.get("layout")}) or ["Layout A", "Layout B", "Layout C"]
    for c, lay in enumerate(["Scenario ID"] + layouts, 1):
        cell = ws5.cell(row=gr, column=c, value=lay)
        cell.font = BOLD
        cell.fill = HEAD_FILL
    for i, (sid, _txt) in enumerate(SCENARIOS):
        rr = gr + 1 + i
        ws5.cell(row=rr, column=1, value=sid).font = FONT
        for c, lay in enumerate(layouts, 2):
            cell = ws5.cell(row=rr, column=c, value=f'=COUNTIFS({rng("Scenario ID")},$A{rr},{rng("Scene Layout")},"{lay}")')
            cell.font = FONT
    ws5.freeze_panes = "A5"
    # ── Definitions ─────────────────────────────────────────────────────
    ws6 = wb.create_sheet("Definitions", 5)
    header(ws6, 1, ["Column / item", "Meaning"], [40, 120])
    for r, (k, v) in enumerate(DEFINITIONS, 2):
        ws6.cell(row=r, column=1, value=k).font = BOLD
        c2 = ws6.cell(row=r, column=2, value=v)
        c2.font = FONT
        c2.alignment = WRAP
    # ── _trials (hidden) ────────────────────────────────────────────────
    ws7 = wb.create_sheet("_trials")
    ws7["A1"] = "trial folder"
    for r, rec in enumerate(recs, 2):
        ws7.cell(row=r, column=1, value=rec["folder"])
    ws7.sheet_state = "hidden"
    wb.active = 0
    for w in wb.worksheets:
        w.sheet_view.tabSelected = (w.title == "Experiments")
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    tmp = path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, path)
    return {"path": path, "trials": n, "conversation_lines": sum(len(r["conversation"]) for r in recs),
            "instructions": sum(len(r["instructions"]) for r in recs)}


def recalc(path: str, timeout: int = 60) -> dict:
    """Store computed values for every formula (LibreOffice, private profile),
    so viewers that read cached values show numbers, then count errors."""
    import shutil
    import subprocess
    soffice = shutil.which("soffice")
    if not soffice:
        return {"recalculated": False, "why": "LibreOffice (soffice) not found — Excel computes on opening"}
    profile = Path(os.path.expanduser("~/.cache/morty_xlsx_profile"))
    macro_dir = profile / "user" / "basic" / "Standard"
    env = {**os.environ, "SAL_USE_VCLPLUGIN": "svp"}
    if not (macro_dir / "Module1.xba").exists():
        profile.mkdir(parents=True, exist_ok=True)
        subprocess.run([soffice, "--headless", "--terminate_after_init", f"-env:UserInstallation={profile.as_uri()}"],
                       capture_output=True, timeout=timeout, env=env)
        if not macro_dir.exists():
            return {"recalculated": False, "why": "LibreOffice did not create its profile"}
        (macro_dir / "Module1.xba").write_text(
            '<?xml version="1.0" encoding="UTF-8"?>\n<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">\n'
            '<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">\n'
            "Sub RecalculateAndSave()\n  ThisComponent.calculateAll()\n  ThisComponent.store()\n  ThisComponent.close(True)\nEnd Sub\n"
            "</script:module>")
    # 2026-09-28: recalculate a COPY and swap it in — with the workbook open
    # in LibreOffice (its lock file present) the headless instance waited on
    # the "document in use" question until the timeout.
    work = os.path.join(os.path.dirname(os.path.abspath(path)), ".recalc_" + os.path.basename(path))
    shutil.copy2(path, work)
    before = os.stat(work).st_mtime_ns
    try:
        res = subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation={profile.as_uri()}",
                              "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
                              work], capture_output=True, text=True, timeout=timeout, env=env)
    except subprocess.TimeoutExpired:
        subprocess.run(["pkill", "-f", "[U]serInstallation=" + profile.as_uri()], capture_output=True)
        os.path.exists(work) and os.remove(work)
        return {"recalculated": False, "why": f"LibreOffice timed out after {timeout} s"}
    if res.returncode != 0 or os.stat(work).st_mtime_ns == before:
        os.path.exists(work) and os.remove(work)
        return {"recalculated": False, "why": (res.stderr or "").strip()[:200] or "LibreOffice did not rewrite the file"}
    os.replace(work, path)
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") and c.value.rstrip("!?0/").upper() in (
                        "#VALUE", "#DIV", "#REF", "#NAME", "#NULL", "#NUM", "#N/A"):
                    errors.append(f"{ws.title}!{c.coordinate}={c.value}")
    wb.close()
    return {"recalculated": True, "formula_errors": len(errors), "errors": errors[:20]}


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    out = DEFAULT_XLSX
    if "--out" in argv:
        i = argv.index("--out")
        out = argv[i + 1]
        argv = argv[:i] + argv[i + 2:]
    if not argv:
        print(__doc__)
        return 2
    folders = [str(p) for p in all_trial_dirs()] if argv == ["--all"] else [str(Path(a).expanduser()) for a in argv]
    res = build(folders, out)
    res.update(recalc(out))
    print(json.dumps(res, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
